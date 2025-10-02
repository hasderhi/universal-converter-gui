from .base import Converter
import os
import tempfile
from pathlib import Path
from openpyxl import load_workbook, Workbook
import xlrd

class ExcelConverter(Converter):
    name = "Excel Converter"
    input_formats = ["xls", "xlsx"]
    output_formats = ["xls", "xlsx"]
    info = "At this moment, Excel conversions only preserve cell values (not styles or formulas)."

    def convert(self, input_file: str, output_format: str, options: dict) -> str:
        input_file = Path(input_file)  # Ensure it's a Path object
        base = input_file.stem
        tmp_out = Path(tempfile.gettempdir()) / f"{base}.{output_format}"

        if input_file.suffix.lower() == ".xls" and output_format.lower() == "xlsx":
            # Convert XLS > XLSX
            book = xlrd.open_workbook(str(input_file))
            wb = Workbook()
            ws = wb.active

            sheet = book.sheet_by_index(0)
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    ws.cell(row=row+1, column=col+1, value=sheet.cell_value(row, col))

            wb.save(tmp_out)

        elif input_file.suffix.lower() == ".xlsx" and output_format.lower() == "xls":
            # Convert XLSX > XLS
            wb = load_workbook(str(input_file))
            ws = wb.active

            import xlwt
            book = xlwt.Workbook()
            sheet = book.add_sheet("Sheet1")

            for i, row in enumerate(ws.iter_rows(values_only=True)):
                for j, val in enumerate(row):
                    sheet.write(i, j, val)

            book.save(str(tmp_out))

        else:
            raise ValueError(f"Unsupported conversion {input_file} -> {output_format}")

        return str(tmp_out)
